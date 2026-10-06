import sys
sys.path.insert(0, sys.argv[1])
import openpyxl
wb = openpyxl.load_workbook(sys.argv[2], read_only=True, data_only=True)
for ws in wb.worksheets:
    print('SHEET', ws.title, ws.max_row, ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= int(sys.argv[3]): break
        print(i, row[:int(sys.argv[4])])
