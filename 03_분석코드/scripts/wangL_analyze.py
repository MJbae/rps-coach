import sys, statistics as st
sys.path.insert(0, '/private/tmp/claude-501/-Users-baemanjin-orca-projects-----/9282a235-4fc4-4c50-add0-37a1d5aaa853/scratchpad/pylib')
import openpyxl
BEATS = {'R': 'S', 'P': 'R', 'S': 'P'}       # key beats value
UP = {'R': 'P', 'P': 'S', 'S': 'R'}          # hand that beats key ("upgrade")
DOWN = {'R': 'S', 'P': 'R', 'S': 'P'}        # hand that key beats ("downgrade")
def outcome(me, opp):
    if me == opp: return 'T'
    return 'W' if BEATS[me] == opp else 'L'
players = []
for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(rows) if r and isinstance(r[0], str) and 'Round number' in r[0])
        hdr = rows[hdr_i]
        pcol = next(j for j, h in enumerate(hdr) if isinstance(h, str) and h.strip().startswith('Player') and 'played' in h)
        acol = next(j for j, h in enumerate(hdr) if isinstance(h, str) and 'Multi' in h and 'played' in h)
        seq = []
        for r in rows[hdr_i + 1:]:
            if r is None or r[0] is None: continue
            p, a = r[pcol], r[acol]
            if p in ('R', 'P', 'S') and a in ('R', 'P', 'S'):
                seq.append((p, a))
        players.append((name, path.split('_')[-2], seq))
print('players parsed:', len(players), 'rounds per player (min,max):', min(len(s) for _, _, s in players), max(len(s) for _, _, s in players))
# AI performance
nets, wr, lr = [], [], []
for name, f, seq in players:
    w = sum(1 for p, a in seq if outcome(a, p) == 'W'); l = sum(1 for p, a in seq if outcome(a, p) == 'L'); n = len(seq)
    nets.append(w - l); wr.append(w / n); lr.append(l / n)
print('AI net (wins-losses) per player: mean %.2f, median %.1f, n=%d, players with AI net>0: %d' % (st.mean(nets), st.median(nets), len(nets), sum(1 for x in nets if x > 0)))
print('AI win rate mean %.4f, loss rate mean %.4f, diff %.4f' % (st.mean(wr), st.mean(lr), st.mean(wr) - st.mean(lr)))
for tag in ('MOESM1', 'MOESM2'):
    sub = [w - l for (name, f, seq), w, l in zip(players, [sum(1 for p, a in s if outcome(a, p) == 'W') for _, _, s in players], [sum(1 for p, a in s if outcome(a, p) == 'L') for _, _, s in players]) if f == tag]
    if sub: print(' ', tag, 'n=%d mean AI net %.2f sd %.2f' % (len(sub), st.mean(sub), st.stdev(sub)))
# human marginal and conditional transitions (pooled over players, and mean of per-player rates)
from collections import Counter
mar = Counter(); cond = {o: Counter() for o in 'WTL'}
per_player = {o: {'stay': [], 'up': [], 'down': []} for o in 'WTL'}
for name, f, seq in players:
    pc = {o: Counter() for o in 'WTL'}
    for p, a in seq: mar[p] += 1
    for (p0, a0), (p1, a1) in zip(seq, seq[1:]):
        o = outcome(p0, a0)
        t = 'stay' if p1 == p0 else ('up' if p1 == UP[p0] else 'down')
        cond[o][t] += 1; pc[o][t] += 1
    for o in 'WTL':
        tot = sum(pc[o].values())
        if tot >= 10:
            for t in ('stay', 'up', 'down'): per_player[o][t].append(pc[o][t] / tot)
N = sum(mar.values())
print('human marginals: R %.4f P %.4f S %.4f (N=%d)' % (mar['R'] / N, mar['P'] / N, mar['S'] / N, N))
for o in 'WTL':
    tot = sum(cond[o].values())
    print('after human %s (pooled n=%d): stay %.3f up %.3f down %.3f | mean-of-players (n=%d): stay %.3f up %.3f down %.3f' % (
        o, tot, cond[o]['stay'] / tot, cond[o]['up'] / tot, cond[o]['down'] / tot, len(per_player[o]['stay']),
        st.mean(per_player[o]['stay']), st.mean(per_player[o]['up']), st.mean(per_player[o]['down'])))
