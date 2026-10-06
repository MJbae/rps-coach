import sys
sys.path.insert(0, '/private/tmp/claude-501/-Users-baemanjin-orca-projects-----/9282a235-4fc4-4c50-add0-37a1d5aaa853/scratchpad/pylib')
import openpyxl
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
nrows = int(sys.argv[2]) if len(sys.argv) > 2 else 8
print('sheets:', len(wb.sheetnames), wb.sheetnames[:60])
for name in wb.sheetnames[:3]:
    ws = wb[name]
    print('=== SHEET', name, ws.max_row, 'x', ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= nrows: break
        print(i, [c for c in row[:30]])
